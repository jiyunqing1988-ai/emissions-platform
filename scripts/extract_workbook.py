from __future__ import annotations

import json
import math
import os
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_SOURCE = (
    "/Users/jyq/Library/Containers/com.tencent.xinWeChat/Data/Documents/"
    "xwechat_files/wxid_nllecx4jflad22_8737/temp/drag/"
    "中国企业温室气体排放核算与报告因子库-20260726.xlsx"
)

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
ASSET_DIR = ROOT / "assets"
DOWNLOAD_DIR = ROOT / "downloads"


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        return round(value, 12)
    return str(value).strip() if not isinstance(value, (int, float)) else value


def text(value) -> str:
    value = clean(value)
    return "" if value == "" else str(value)


def trim_row(values):
    values = [clean(v) for v in values]
    while values and values[-1] == "":
        values.pop()
    return values


def actual_max_col(ws):
    max_col = 0
    for row in ws.iter_rows():
        for cell in row:
            if cell.value not in (None, ""):
                max_col = max(max_col, cell.column)
    return max_col


def sheet_rows(ws):
    max_col = actual_max_col(ws)
    rows = []
    for row in ws.iter_rows(min_col=1, max_col=max_col):
        values = trim_row([cell.value for cell in row])
        if any(v != "" for v in values):
            rows.append({"row": row[0].row, "cells": values})
    return rows


def find_header_row(ws):
    required_groups = [
        {"排放源", "活动数据单位"},
        {"类别名称", "活动数据单位"},
        {"排放源类别", "活动数据单位"},
        {"类别", "类别名称", "活动数据单位"},
    ]
    scan_rows = min(ws.max_row, 20)
    scan_cols = min(actual_max_col(ws), 40)
    for row_idx in range(1, scan_rows + 1):
        labels = {text(ws.cell(row_idx, col).value) for col in range(1, scan_cols + 1)}
        labels.discard("")
        if any(group.issubset(labels) for group in required_groups):
            return row_idx
    return 11


def header_map(ws, header_row):
    max_col = actual_max_col(ws)
    headers = []
    for col in range(1, max_col + 1):
        label = clean(ws.cell(header_row, col).value)
        if label != "":
            headers.append((col, str(label)))
    return headers


def formula_display(value, headers):
    if not (isinstance(value, str) and value.startswith("=")):
        return value
    display = value
    display = re.sub(r"GWP!\$?E\$?6", "GWP(CH4)", display)
    display = re.sub(r"GWP!\$?E\$?9", "GWP(N2O)", display)
    col_labels = {
        col: label
        .replace("kg ", "")
        .replace("t ", "")
        .replace("₂", "2")
        .replace("₄", "4")
        .replace("₂", "2")
        .replace("CO2e", "CO2e")
        for col, label in headers
    }

    def repl(match):
        col_letters = match.group(1)
        col_number = 0
        for char in col_letters:
            col_number = col_number * 26 + ord(char.upper()) - ord("A") + 1
        return col_labels.get(col_number, match.group(0))

    display = re.sub(r"(?<![A-Za-z!])\$?([A-Z]{1,3})\$?\d+", repl, display)
    return display


def first_header(fields, candidates):
    for candidate in candidates:
        if candidate in fields:
            return candidate
    return ""


def first_present(fields, names):
    for name in names:
        value = fields.get(name)
        if value not in (None, ""):
            return value
    return ""


def extract_factors(values_wb, formula_wb):
    specs = [
        ("范围一", "Scope 1", "S1"),
        ("范围二", "Scope 2", "S2"),
        ("范围三", "Scope 3", "S3"),
    ]
    records = []
    for sheet_name, scope_code, prefix in specs:
        ws_values = values_wb[sheet_name]
        ws_formula = formula_wb[sheet_name]
        header_row = find_header_row(ws_values)
        headers = header_map(ws_values, header_row)
        index = 0
        for row_idx in range(header_row + 1, ws_values.max_row + 1):
            fields = {}
            formulas = {}
            has_value = False
            for col, label in headers:
                value = clean(ws_values.cell(row_idx, col).value)
                formula_value = clean(ws_formula.cell(row_idx, col).value)
                fields[label] = value
                if isinstance(formula_value, str) and formula_value.startswith("="):
                    formulas[label] = formula_display(formula_value, headers)
                if value != "":
                    has_value = True
            if not has_value:
                continue

            index += 1
            if sheet_name == "范围三":
                category = first_present(fields, ["类别名称", "排放源类别"])
            else:
                category = first_present(fields, ["排放源", "类别名称", "排放源类别"])
            source_name = first_present(fields, ["燃料/物料/过程", "能源类型", "产品或服务", "排放源"])
            if sheet_name == "范围二":
                feature = first_present(fields, ["区域", "系统边界"])
            else:
                feature = first_present(fields, ["排放源排放特征", "系统边界", "区域"])
            unit = first_present(fields, ["活动数据单位"])
            notes = first_present(fields, ["备注"])
            factor_header = first_header(fields, ["kg CO2e", "t CO2e", "kg CO₂", "t CO₂"])
            factor_value = fields.get(factor_header, "")
            gas_values = {
                header: value
                for header, value in fields.items()
                if header in {"kg CO2e", "t CO2e", "kg CO₂", "t CO₂", "kg CH₄", "t CH₄", "kg N₂O", "t N₂O"}
                and value != ""
            }
            supporting_values = {
                header: value
                for header, value in fields.items()
                if header in {"计算方法", "数据来源", "系统边界"}
                if fields.get(header, "") != ""
            }
            search_parts = [
                sheet_name,
                scope_code,
                str(category),
                str(source_name),
                str(feature),
                str(unit),
                str(notes),
                " ".join(f"{k} {v}" for k, v in gas_values.items()),
            ]
            records.append(
                {
                    "id": f"{prefix}-{index:03d}",
                    "scope": sheet_name,
                    "scopeCode": scope_code,
                    "category": category,
                    "sourceName": source_name,
                    "feature": feature,
                    "activityUnit": unit,
                    "factorHeader": factor_header,
                    "factorValue": factor_value,
                    "factorUnit": f"{factor_header}/{unit}" if unit else factor_header,
                    "gasValues": gas_values,
                    "supportingValues": supporting_values,
                    "otherGhg": fields.get("其他温室气体类别", ""),
                    "notes": notes,
                    "row": row_idx,
                    "fields": fields,
                    "formulas": formulas,
                    "searchText": " ".join(search_parts).lower(),
                }
            )
    return records


def parse_references(rows):
    references = []
    pattern = re.compile(r"^\[(\d+)\]\s*(.*)$")
    for row in rows:
        joined = " ".join(str(cell) for cell in row["cells"] if cell != "").strip()
        match = pattern.match(joined)
        if match:
            references.append({"id": match.group(1), "text": match.group(2)})
        elif joined and joined != "参考文献":
            references.append({"id": str(len(references) + 1), "text": joined})
    return references


def parse_qa(rows):
    items = []
    current_section = ""
    current = None
    for row in rows:
        cells = [str(cell).strip() for cell in row["cells"] if str(cell).strip()]
        if not cells:
            continue
        line = cells[0]
        if line.startswith("Q"):
            current = {"section": current_section, "question": line, "answer": ""}
            items.append(current)
        elif current and not current["answer"]:
            current["answer"] = line
        elif len(line) <= 18 and not line.startswith("范围") and not line.startswith("常见"):
            current_section = line
    return items


def main():
    source = Path(os.environ.get("SOURCE_XLSX", DEFAULT_SOURCE))
    if not source.exists():
        raise FileNotFoundError(source)
    version_match = re.search(r"(\d{8}(?:-[A-Za-z0-9]+)?)", source.stem)
    version = version_match.group(1) if version_match else source.stem

    DATA_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    DOWNLOAD_DIR.mkdir(parents=True, exist_ok=True)

    values_wb = load_workbook(source, data_only=True)
    formula_wb = load_workbook(source, data_only=False)
    records = extract_factors(values_wb, formula_wb)

    content_sheet_names = ["基本信息", "导航目录", "Q&A", "参考文献", "参考文献-nzy"]
    pages = {}
    for name in content_sheet_names:
        if name not in values_wb.sheetnames:
            continue
        rows = sheet_rows(values_wb[name])
        pages[name] = {"name": name, "rows": rows}

    references = parse_references(pages["参考文献"]["rows"])
    qa_items = parse_qa(pages["Q&A"]["rows"])
    categories = {}
    for record in records:
        key = f"{record['scope']}::{record['category']}"
        categories[key] = categories.get(key, 0) + 1

    payload = {
        "metadata": {
            "title": "中国企业温室气体排放核算与报告因子库",
            "version": version,
            "applicablePeriod": "2025年1月1日~2025年12月31日",
            "sourceFile": source.name,
            "totalFactors": len(records),
            "scopeCounts": {
                "范围一": sum(1 for item in records if item["scope"] == "范围一"),
                "范围二": sum(1 for item in records if item["scope"] == "范围二"),
                "范围三": sum(1 for item in records if item["scope"] == "范围三"),
            },
            "downloadFile": f"downloads/{source.name}",
        },
        "factors": records,
        "contentPages": pages,
        "references": references,
        "qaItems": qa_items,
        "categoryCounts": categories,
    }
    (DATA_DIR / "platform-data.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    download_target = ROOT / payload["metadata"]["downloadFile"]
    if source.resolve() != download_target.resolve():
        shutil.copyfile(source, download_target)
    print(json.dumps(payload["metadata"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
